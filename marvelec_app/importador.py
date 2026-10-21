# marvelec_app/importador.py
"""
Carga del listado de puntos desde Excel (historia de usuario 8).

Acepta dos formatos, detectados por hoja:

1. Planilla de paquetes de MARVELEC (la que se arma tras medir): una hoja por rack
   ("RACK_1H" o una celda "RACK 1H"), una columna por paquete con "PAQUETE 01", los códigos
   (DT-001, VD-009, RI-040...) hacia abajo y al final "36 METROS". Las tablas de totales se ignoran.

2. Tabla simple con encabezados: rack | paquete | codigo | zona | metros.

La carga nunca borra puntos: crea los nuevos y actualiza paquete/metraje de los existentes.
"""
import re
import unicodedata
from dataclasses import asdict, dataclass

from django.db import transaction

from .models import Paquete, Punto, Rack, TipoPunto, Zona

RE_RACK = re.compile(r'^\s*RACK[\s_\-]*([A-Z0-9][A-Z0-9\-]*)\s*$', re.I)
RE_PAQUETE = re.compile(r'^\s*(?:PAQUETE|PQ)\s*\.?\s*0*(\d+)\s*$', re.I)
RE_METROS = re.compile(r'(\d+(?:[.,]\d+)?)\s*(?:METROS?|MTS?\.?|M)\s*$', re.I)
RE_CANTIDAD = re.compile(r'^\s*\d+\s*PUNTOS?\s*$', re.I)
RE_CODIGO = re.compile(r'^([A-Z]{1,6})-?(\d{1,5})$')

ENCABEZADOS = {
    'rack': 'rack', 'paquete': 'paquete', 'pq': 'paquete', 'codigo': 'codigo', 'punto': 'codigo',
    'zona': 'zona', 'sector': 'zona', 'metros': 'metros', 'metraje': 'metros', 'largo': 'metros',
}
MAX_FILAS_VACIAS = 6


@dataclass
class FilaPunto:
    rack: str
    sigla: str
    numero: int
    paquete: int = None
    zona: str = ''
    metros: int = None
    origen: str = ''


class Resultado:
    def __init__(self):
        self.puntos = []
        self.errores = []
        self.advertencias = []

    def resumen_por_rack(self):
        """[{'rack', 'puntos', 'paquetes', 'tipos': {sigla: n}}] para la vista previa."""
        racks = {}
        for f in self.puntos:
            r = racks.setdefault(f.rack, {'rack': f.rack, 'puntos': 0, 'paquetes': set(), 'tipos': {}})
            r['puntos'] += 1
            if f.paquete is not None:
                r['paquetes'].add(f.paquete)
            r['tipos'][f.sigla] = r['tipos'].get(f.sigla, 0) + 1
        return [
            {**r, 'paquetes': len(r['paquetes']), 'tipos': sorted(r['tipos'].items())}
            for r in sorted(racks.values(), key=lambda r: r['rack'])
        ]

    def serializar(self):
        return [asdict(f) for f in self.puntos]

    @staticmethod
    def desde_lista(lista):
        resultado = Resultado()
        resultado.puntos = [FilaPunto(**f) for f in lista]
        return resultado


def _texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _normalizar(texto):
    sin_tildes = unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode()
    return sin_tildes.strip().lower()


def leer_codigo(valor):
    """'DT-015', 'dt 15', 'DT`-042' (error de tipeo real de la planilla) -> ('DT', 15)."""
    limpio = re.sub(r'[^A-Z0-9]', '', _texto(valor).upper())
    coincidencia = RE_CODIGO.match(limpio)
    if not coincidencia:
        return None
    return coincidencia.group(1), int(coincidencia.group(2))


def _metros(valor):
    texto = _texto(valor)
    if not texto:
        return None
    m = RE_METROS.search(texto)
    numero = m.group(1) if m else texto
    try:
        return round(float(numero.replace(',', '.')))
    except ValueError:
        return None


def _celda(fila, columna):
    return f'{_columna_letra(columna)}{fila}'


def _columna_letra(n):
    letras = ''
    while n:
        n, resto = divmod(n - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


# ---------------------------------------------------------
# Lectura
# ---------------------------------------------------------

def leer_excel(archivo):
    import openpyxl
    try:
        libro = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    except Exception:
        resultado = Resultado()
        resultado.errores.append('No se pudo leer el archivo. Debe ser una planilla Excel (.xlsx).')
        return resultado

    resultado = Resultado()
    for hoja in libro.worksheets:
        filas = [list(f) for f in hoja.iter_rows(values_only=True)]
        if not any(any(_texto(c) for c in f) for f in filas):
            continue
        encabezado = _buscar_encabezado(filas)
        if encabezado is not None:
            _leer_tabla(hoja.title, filas, encabezado, resultado)
        else:
            _leer_planilla_paquetes(hoja.title, filas, resultado)
    libro.close()

    _quitar_duplicados(resultado)
    if not resultado.puntos and not resultado.errores:
        resultado.errores.append('No se encontraron códigos de puntos (ej. DT-001) en el archivo.')
    return resultado


def _buscar_encabezado(filas):
    for i, fila in enumerate(filas[:10]):
        columnas = {}
        for j, valor in enumerate(fila):
            clave = ENCABEZADOS.get(_normalizar(_texto(valor)))
            if clave and clave not in columnas:
                columnas[clave] = j
        if 'codigo' in columnas:
            return i, columnas
    return None


def _leer_tabla(nombre_hoja, filas, encabezado, resultado):
    inicio, columnas = encabezado
    rack_hoja = _rack_de_hoja(nombre_hoja, filas)

    def valor(fila, clave):
        j = columnas.get(clave)
        return fila[j] if j is not None and j < len(fila) else None

    for i, fila in enumerate(filas[inicio + 1:], start=inicio + 2):
        texto_codigo = _texto(valor(fila, 'codigo'))
        if not texto_codigo:
            continue
        codigo = leer_codigo(texto_codigo)
        if not codigo:
            resultado.advertencias.append(f'Hoja "{nombre_hoja}", fila {i}: "{texto_codigo}" no es un código válido.')
            continue
        rack = _texto(valor(fila, 'rack')).upper().removeprefix('RACK').strip(' _-') or rack_hoja
        if not rack:
            resultado.errores.append(f'Hoja "{nombre_hoja}", fila {i}: falta el rack.')
            continue
        paquete_txt = _texto(valor(fila, 'paquete'))
        paquete = None
        if paquete_txt:
            m = RE_PAQUETE.match(paquete_txt) or re.match(r'^\s*0*(\d+)\s*$', paquete_txt)
            paquete = int(m.group(1)) if m else None
            if paquete is None:
                resultado.advertencias.append(f'Hoja "{nombre_hoja}", fila {i}: paquete "{paquete_txt}" no reconocido.')
        resultado.puntos.append(FilaPunto(
            rack=rack, sigla=codigo[0], numero=codigo[1], paquete=paquete,
            zona=_texto(valor(fila, 'zona')), metros=_metros(valor(fila, 'metros')),
            origen=f'{nombre_hoja}!fila {i}',
        ))


def _rack_de_hoja(nombre_hoja, filas):
    for fila in filas[:15]:
        for valor in fila:
            m = RE_RACK.match(_texto(valor))
            if m:
                return m.group(1).upper()
    m = RE_RACK.match(nombre_hoja.replace('_', ' '))
    return m.group(1).upper() if m else ''


def _leer_planilla_paquetes(nombre_hoja, filas, resultado):
    rack = _rack_de_hoja(nombre_hoja, filas)
    cabeceras = [
        (i, j, int(RE_PAQUETE.match(_texto(v)).group(1)))
        for i, fila in enumerate(filas) for j, v in enumerate(fila)
        if RE_PAQUETE.match(_texto(v))
    ]
    if not cabeceras:
        hay_codigos = any(leer_codigo(v) for f in filas for v in f)
        if hay_codigos:
            resultado.advertencias.append(
                f'Hoja "{nombre_hoja}": hay códigos pero no encabezados "PAQUETE NN" ni columna "codigo"; se ignoró.'
            )
        return
    if not rack:
        resultado.errores.append(
            f'Hoja "{nombre_hoja}": no se encontró el nombre del rack (hoja "RACK_1H" o celda "RACK 1H").'
        )
        return

    for fila_cab, col, paquete in cabeceras:
        encontrados, metros, esperados, vacias = [], None, None, 0
        for i in range(fila_cab + 1, len(filas)):
            texto = _texto(filas[i][col]) if col < len(filas[i]) else ''
            if not texto:
                vacias += 1
                if encontrados and vacias > MAX_FILAS_VACIAS:
                    break
                continue
            vacias = 0
            if RE_PAQUETE.match(texto):
                break
            if RE_CANTIDAD.match(texto):
                esperados = int(re.match(r'\s*(\d+)', texto).group(1))
                continue
            codigo = leer_codigo(texto)
            if codigo:
                encontrados.append((codigo, _celda(i + 1, col + 1)))
                continue
            if RE_METROS.search(texto):
                metros = _metros(texto)
                break
            resultado.advertencias.append(
                f'Hoja "{nombre_hoja}", celda {_celda(i + 1, col + 1)}: "{texto}" no es un código válido.'
            )

        if esperados is not None and esperados != len(encontrados):
            resultado.advertencias.append(
                f'Rack {rack}, paquete {paquete:02d}: dice {esperados} puntos pero se encontraron {len(encontrados)}.'
            )
        for (sigla, numero), celda in encontrados:
            resultado.puntos.append(FilaPunto(
                rack=rack, sigla=sigla, numero=numero, paquete=paquete, metros=metros,
                origen=f'{nombre_hoja}!{celda}',
            ))


def _quitar_duplicados(resultado):
    vistos = {}
    unicos = []
    for fila in resultado.puntos:
        clave = (fila.rack, fila.sigla, fila.numero)
        if clave in vistos:
            resultado.advertencias.append(
                f'{fila.sigla}-{fila.numero:03d} del rack {fila.rack} aparece dos veces '
                f'({vistos[clave]} y {fila.origen}); se usó la primera.'
            )
            continue
        vistos[clave] = fila.origen
        unicos.append(fila)
    resultado.puntos = unicos


# ---------------------------------------------------------
# Vista previa y aplicación
# ---------------------------------------------------------

def comparar(resultado, obra):
    """Cuántos puntos son nuevos y cuántos ya existían en la obra."""
    existentes = set(
        Punto.objects.filter(rack__obra=obra).values_list('rack__nombre', 'tipo__sigla', 'numero')
    )
    nuevos = sum(1 for f in resultado.puntos if (f.rack, f.sigla, f.numero) not in existentes)
    racks_existentes = set(Rack.objects.filter(obra=obra).values_list('nombre', flat=True))
    racks_nuevos = sorted({f.rack for f in resultado.puntos} - racks_existentes)
    return {'nuevos': nuevos, 'actualizados': len(resultado.puntos) - nuevos, 'racks_nuevos': racks_nuevos}


@transaction.atomic
def aplicar(resultado, obra):
    """Crea o actualiza racks, paquetes, zonas, tipos y puntos. Devuelve un resumen."""
    tipos = {t.sigla: t for t in TipoPunto.objects.all()}
    racks = {r.nombre: r for r in Rack.objects.filter(obra=obra)}
    zonas = {z.nombre.lower(): z for z in Zona.objects.filter(obra=obra)}
    paquetes = {(p.rack_id, p.numero): p for p in Paquete.objects.filter(rack__obra=obra)}
    vistos = set()  # paquetes cuyo metraje ya se fijó en esta carga
    creados = actualizados = 0

    for fila in resultado.puntos:
        tipo = tipos.get(fila.sigla)
        if tipo is None:
            tipo = tipos[fila.sigla] = TipoPunto.objects.create(sigla=fila.sigla, orden=len(tipos) + 1)
        rack = racks.get(fila.rack)
        if rack is None:
            rack = racks[fila.rack] = Rack.objects.create(obra=obra, nombre=fila.rack)

        zona = None
        if fila.zona:
            zona = zonas.get(fila.zona.lower())
            if zona is None:
                zona = zonas[fila.zona.lower()] = Zona.objects.create(obra=obra, nombre=fila.zona)

        paquete = None
        if fila.paquete is not None:
            clave = (rack.pk, fila.paquete)
            paquete = paquetes.get(clave)
            if paquete is None:
                paquete = paquetes[clave] = Paquete.objects.create(
                    rack=rack, numero=fila.paquete, metros=fila.metros, zona=zona,
                )
            else:
                cambios = []
                if clave not in vistos and fila.metros is not None and paquete.metros != fila.metros:
                    paquete.metros = fila.metros
                    cambios.append('metros')
                if zona and paquete.zona_id != zona.pk:
                    paquete.zona = zona
                    cambios.append('zona')
                if cambios:
                    paquete.save(update_fields=cambios)
            vistos.add(clave)

        # El metraje queda en el punto sólo si difiere del de su paquete (tabla simple con largos por punto).
        metros_punto = fila.metros if paquete is None or fila.metros != paquete.metros else None
        _, creado = Punto.objects.update_or_create(
            rack=rack, tipo=tipo, numero=fila.numero, defaults={'paquete': paquete, 'metros': metros_punto},
        )
        creados += creado
        actualizados += not creado

    return {'creados': creados, 'actualizados': actualizados, 'racks': len({f.rack for f in resultado.puntos})}


def plantilla_tabla():
    """Libro Excel de ejemplo para el formato de tabla simple."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Puntos'
    hoja.append(['rack', 'paquete', 'codigo', 'zona', 'metros'])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='3A288F')
    ejemplos = [
        ('1H', 1, 'VD-001', 'Sótano 1 - Ala Norte', 36),
        ('1H', 1, 'RI-001', 'Sótano 1 - Ala Norte', 36),
        ('1H', 2, 'DT-001', 'Sótano 1 - Ala Norte', 60),
        ('1H', 2, 'DT-002', 'Sótano 1 - Ala Norte', 60),
    ]
    for fila in ejemplos:
        hoja.append(fila)
    for col, ancho in zip('ABCDE', (10, 10, 12, 28, 10)):
        hoja.column_dimensions[col].width = ancho
    return libro
