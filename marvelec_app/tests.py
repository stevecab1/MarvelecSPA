import io
import json
import shutil
import tempfile
import uuid
from datetime import timedelta
from unittest import mock

import openpyxl
from django.contrib.auth.models import User
from django.core import signing
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from . import consultas, etapas, importador
from .models import (
    Asistencia, AvancePunto, Incidencia, Notificacion, Obra, Paquete, PerfilTrabajador, Punto, Rack, ReporteAvance,
    ResumenDiarioEnviado, SuscripcionPush, TipoIncidencia, TipoPunto, Zona,
)
from .servicios import parsear_codigos, resolver_codigos

MEDIA_PRUEBAS = tempfile.mkdtemp(prefix='marvelec-tests-')


def tearDownModule():
    shutil.rmtree(MEDIA_PRUEBAS, ignore_errors=True)


def imagen(nombre='foto.png'):
    from PIL import Image
    buffer = io.BytesIO()
    Image.new('RGB', (4, 4), 'purple').save(buffer, 'PNG')
    return SimpleUploadedFile(nombre, buffer.getvalue(), content_type='image/png')


# =========================================================
# Escenario común: dos obras, racks con puntos, trabajadores, supervisores y gerencia
# =========================================================

@override_settings(MEDIA_ROOT=MEDIA_PRUEBAS)
class BaseEscenario(TestCase):

    def setUp(self):
        self.obra_a = Obra.objects.create(nombre="Hospital A")
        self.obra_b = Obra.objects.create(nombre="Hospital B")
        self.zona_a = Zona.objects.create(obra=self.obra_a, nombre="Sótano 1")
        self.rack_a = Rack.objects.create(obra=self.obra_a, nombre="1H")
        self.rack_b = Rack.objects.create(obra=self.obra_b, nombre="2A")
        self.dt = TipoPunto.objects.get(sigla='DT')
        self.vd = TipoPunto.objects.get(sigla='VD')
        self.paq1 = Paquete.objects.create(rack=self.rack_a, numero=1, zona=self.zona_a, metros=36)
        self.paq2 = Paquete.objects.create(rack=self.rack_a, numero=2, metros=60)
        for n in range(1, 7):
            Punto.objects.create(rack=self.rack_a, paquete=self.paq1 if n <= 3 else self.paq2, tipo=self.dt, numero=n)
        Punto.objects.create(rack=self.rack_a, paquete=self.paq2, tipo=self.vd, numero=1)
        for n in range(1, 4):
            Punto.objects.create(rack=self.rack_b, tipo=self.dt, numero=n)

        self.trab_a = self._usuario("trab_a", PerfilTrabajador.ROL_TRABAJADOR, self.obra_a, "Pedro", "Soto")
        self.trab_a2 = self._usuario("trab_a2", PerfilTrabajador.ROL_TRABAJADOR, self.obra_a, "Juan", "Pérez")
        self.trab_b = self._usuario("trab_b", PerfilTrabajador.ROL_TRABAJADOR, self.obra_b)
        self.sup_a = self._usuario("sup_a", PerfilTrabajador.ROL_SUPERVISOR, self.obra_a)
        self.sup_b = self._usuario("sup_b", PerfilTrabajador.ROL_SUPERVISOR, self.obra_b)
        self.gerente = self._usuario("gerente", PerfilTrabajador.ROL_GERENCIA, None)

    def _usuario(self, username, rol, obra, nombre='', apellido=''):
        u = User.objects.create_user(username=username, password="clave12345", first_name=nombre, last_name=apellido)
        PerfilTrabajador.objects.create(usuario=u, rol=rol, obra_asignada=obra)
        return u

    def _login(self, usuario):
        self.client.login(username=usuario, password="clave12345")

    def _p(self, codigo, rack=None):
        sigla, numero = codigo.split('-')
        return Punto.objects.get(rack=rack or self.rack_a, tipo__sigla=sigla, numero=int(numero))

    def _avance(self, trabajador, codigos, claves, rack=None, **extra):
        """Crea un reporte con avances directamente (sin validar), para preparar escenarios."""
        rack = rack or self.rack_a
        reporte = ReporteAvance.objects.create(trabajador=trabajador, obra=rack.obra, rack=rack, **extra)
        AvancePunto.objects.bulk_create([
            AvancePunto(reporte=reporte, punto=self._p(c, rack), etapa=e) for c in codigos for e in claves
        ])
        reporte.calcular_total_puntos()
        return reporte

    def _post_avance(self, codigos='', puntos=(), claves=('tendido',), rack=None, **extra):
        datos = {
            'rack': (rack or self.rack_a).pk, 'etapas': list(claves), 'puntos': [p.pk for p in puntos],
            'codigos': codigos, 'comentario': 'ok', **extra,
        }
        return self.client.post(reverse('marvelec_app:api_crear_reporte'), datos)


# =========================================================
# Códigos rápidos ("como en WhatsApp")
# =========================================================

class CodigosTests(BaseEscenario):

    def test_guion_lista_puntos_y_a_indica_rango(self):
        pares, errores = parsear_codigos('Dt 1-2-3-17, VD 9; ri 5 a 7 y dt-015')
        self.assertEqual(pares, {('DT', 1), ('DT', 2), ('DT', 3), ('DT', 17), ('VD', 9),
                                 ('RI', 5), ('RI', 6), ('RI', 7), ('DT', 15)})
        self.assertEqual(errores, [])

    def test_numero_sin_tipo_es_error(self):
        pares, errores = parsear_codigos('12 13')
        self.assertFalse(pares)
        self.assertEqual(len(errores), 2)

    def test_resolver_avisa_codigos_inexistentes(self):
        puntos, errores = resolver_codigos(self.rack_a, 'DT 1 a 3, DT 99')
        self.assertEqual([p.codigo for p in puntos], ['DT-001', 'DT-002', 'DT-003'])
        self.assertIn('DT-099', errores[0])


# =========================================================
# Carga de la planilla de paquetes
# =========================================================

@override_settings(MEDIA_ROOT=MEDIA_PRUEBAS)
class ImportadorTests(TestCase):

    def setUp(self):
        self.obra = Obra.objects.create(nombre='Hospital Sótero')

    def _planilla(self):
        """Imita la planilla real: hoja por rack, columnas por paquete, metros al final y totales sueltos."""
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.title = 'RACK_1H'
        hoja['D1'] = 'RACK 1H'
        hoja['A3'], hoja['B3'] = 'PAQUETE 01', 'PAQUETE 02'
        hoja['A5'], hoja['B5'] = '3 PUNTOS', '3 PUNTOS'
        hoja['A7'], hoja['A8'], hoja['A9'], hoja['A10'] = 'VD-001', 'RI-001', 'RI-039', '36 METROS'
        hoja['B7'], hoja['B8'], hoja['B9'] = 'DT-001', 'DT`-042', '60 METROS'  # error de tipeo real
        hoja['D23'], hoja['E23'], hoja['D26'], hoja['E26'] = 'DT', 231, 'TOTAL', 305
        otra = libro.create_sheet('RACK_1A')
        otra['A1'], otra['A2'], otra['A3'] = 'PAQUETE 01', 'DT-001', '20 METROS'
        buffer = io.BytesIO()
        libro.save(buffer)
        buffer.seek(0)
        return buffer

    def test_lee_planilla_de_paquetes(self):
        resultado = importador.leer_excel(self._planilla())
        self.assertEqual(resultado.errores, [])
        codigos = {(f.rack, f.sigla, f.numero, f.paquete, f.metros) for f in resultado.puntos}
        self.assertEqual(codigos, {
            ('1H', 'VD', 1, 1, 36), ('1H', 'RI', 1, 1, 36), ('1H', 'RI', 39, 1, 36),
            ('1H', 'DT', 1, 2, 60), ('1H', 'DT', 42, 2, 60), ('1A', 'DT', 1, 1, 20),
        })
        self.assertTrue(any('dice 3 puntos pero se encontraron 2' in a for a in resultado.advertencias))

    def test_aplicar_crea_y_volver_a_cargar_no_duplica(self):
        resultado = importador.leer_excel(self._planilla())
        resumen = importador.aplicar(resultado, self.obra)
        self.assertEqual(resumen, {'creados': 6, 'actualizados': 0, 'racks': 2})
        self.assertEqual(Rack.objects.filter(obra=self.obra).count(), 2)
        self.assertEqual(Paquete.objects.get(rack__nombre='1H', numero=1).metros, 36)
        self.assertEqual(Punto.objects.get(rack__nombre='1H', tipo__sigla='DT', numero=42).metros_efectivos, 60)

        otra_vez = importador.aplicar(importador.leer_excel(self._planilla()), self.obra)
        self.assertEqual(otra_vez['creados'], 0)
        self.assertEqual(Punto.objects.count(), 6)

    def test_formato_tabla_con_zona_y_metros_por_punto(self):
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.append(['Rack', 'Paquete', 'Código', 'Zona', 'Metros'])
        hoja.append(['1H', 3, 'DT-010', 'Ala Norte', 42])
        hoja.append(['1H', 3, 'DT-011', 'Ala Norte', 45])
        hoja.append(['1H', '', 'XX', '', ''])
        buffer = io.BytesIO()
        libro.save(buffer)
        buffer.seek(0)

        resultado = importador.leer_excel(buffer)
        self.assertEqual(len(resultado.puntos), 2)
        self.assertTrue(resultado.advertencias)  # "XX" no es un código
        importador.aplicar(resultado, self.obra)
        paquete = Paquete.objects.get(rack__nombre='1H', numero=3)
        self.assertEqual((paquete.zona.nombre, paquete.metros), ('Ala Norte', 42))
        self.assertEqual(Punto.objects.get(numero=11).metros_efectivos, 45)

    def test_archivo_que_no_es_excel(self):
        resultado = importador.leer_excel(io.BytesIO(b'no soy un excel'))
        self.assertTrue(resultado.errores)


# =========================================================
# Avance por punto y etapa
# =========================================================

class AvanceTests(BaseEscenario):

    def test_reporte_por_codigos_registra_etapas_y_avisa_a_revisores(self):
        self._login("trab_a")
        r = self._post_avance(codigos='DT 1-2-3')
        self.assertEqual(r.status_code, 201, r.content)
        reporte = ReporteAvance.objects.get()
        self.assertEqual(reporte.total_puntos, 3)
        self.assertEqual(reporte.obra, self.obra_a)
        self.assertEqual(reporte.estado, ReporteAvance.ESTADO_PENDIENTE)
        notificados = set(Notificacion.objects.values_list('usuario__username', flat=True))
        self.assertEqual(notificados, {'sup_a', 'gerente'})

    def test_etapa_sin_requisito_previo_se_rechaza(self):
        self._login("trab_a")
        r = self._post_avance(codigos='DT 1', claves=['bajada'])
        self.assertEqual(r.status_code, 400)
        self.assertIn('falta tendido', r.json()['mensaje'])
        self.assertEqual(ReporteAvance.objects.count(), 0)

    def test_varias_etapas_en_un_reporte_cumplen_sus_requisitos(self):
        self._login("trab_a")
        r = self._post_avance(puntos=[self._p('DT-001'), self._p('DT-002')], claves=['tendido', 'bajada'])
        self.assertEqual(r.status_code, 201)
        self.assertEqual(AvancePunto.objects.count(), 4)

    def test_identificacion_exige_bajada_y_peinado(self):
        self._avance(self.trab_a2, ['DT-001'], ['tendido', 'bajada'])
        self._login("trab_a")
        r = self._post_avance(codigos='DT 1', claves=['identificacion'])
        self.assertEqual(r.status_code, 400)
        self.assertIn('falta peinado de rack', r.json()['mensaje'])
        r = self._post_avance(codigos='DT 1', claves=['peinado', 'identificacion'])
        self.assertEqual(r.status_code, 201)

    def test_etapa_ya_registrada_queda_omitida_con_motivo(self):
        self._avance(self.trab_a2, ['DT-001'], ['tendido'])
        self._login("trab_a")
        r = self._post_avance(codigos='DT 1-2, DT 99')
        self.assertEqual(r.status_code, 201)
        data = r.json()
        self.assertEqual(data['total_puntos'], 1)
        motivos = ' '.join(o['motivo'] for o in data['omitidos'])
        self.assertIn('ya registrada por Juan Pérez', motivos)
        self.assertIn('DT-099', motivos)

    def test_punto_bloqueado_no_avanza(self):
        Incidencia.objects.create(punto=self._p('DT-001'), tipo=TipoIncidencia.objects.first(), reportada_por=self.trab_a2)
        self._login("trab_a")
        r = self._post_avance(codigos='DT 1')
        self.assertEqual(r.status_code, 400)
        self.assertIn('bloqueado', r.json()['mensaje'])

    def test_no_se_reporta_en_rack_de_otra_obra(self):
        self._login("trab_a")
        r = self._post_avance(codigos='DT 1', rack=self.rack_b)
        self.assertEqual(r.status_code, 400)
        self.assertEqual(ReporteAvance.objects.count(), 0)

    def test_api_es_idempotente_con_uuid_cliente(self):
        self._login("trab_a")
        extra = {'uuid_cliente': str(uuid.uuid4()), 'offline': '1'}
        r1 = self._post_avance(codigos='DT 1', **extra)
        r2 = self._post_avance(codigos='DT 1', **extra)
        self.assertEqual((r1.status_code, r2.status_code), (201, 200))
        self.assertTrue(r2.json()['duplicado'])
        self.assertEqual(ReporteAvance.objects.count(), 1)
        self.assertTrue(ReporteAvance.objects.get().enviado_offline)

    def test_api_respeta_hora_del_celular_pero_no_futura(self):
        self._login("trab_a")
        hace_2h = timezone.now() - timedelta(hours=2)
        self._post_avance(codigos='DT 1', fecha_hora_cliente=hace_2h.isoformat())
        self.assertLess(abs((ReporteAvance.objects.get().fecha_hora - hace_2h).total_seconds()), 1)
        self._post_avance(codigos='DT 2', fecha_hora_cliente=(timezone.now() + timedelta(days=2)).isoformat())
        self.assertLess(ReporteAvance.objects.order_by('-pk').first().fecha_hora, timezone.now() + timedelta(minutes=1))

    def test_api_sin_sesion_y_registro_de_otro_usuario(self):
        self.assertEqual(self._post_avance(codigos='DT 1').status_code, 401)
        self._login("trab_a")
        self.assertEqual(self._post_avance(codigos='DT 1', usuario_id=self.trab_a2.pk).status_code, 409)

    def test_foto_del_avance_con_nombre_nuevo_y_antiguo(self):
        """'foto_llegada' lo envían celulares con la versión anterior de la app en caché."""
        self._login("trab_a")
        for numero, clave in ((1, 'foto'), (2, 'foto_llegada')):
            r = self._post_avance(codigos=f'DT {numero}', **{clave: imagen()})
            self.assertEqual(r.status_code, 201, clave)
            self.assertTrue(ReporteAvance.objects.get(pk=r.json()['id']).foto.name.startswith('reportes/avance/'))

    def test_formulario_sin_js(self):
        self._login("trab_a")
        r = self.client.post(reverse('marvelec_app:nuevo_reporte'),
                             {'rack': self.rack_a.pk, 'etapas': ['tendido'], 'codigos': 'DT 4 a 6'})
        self.assertRedirects(r, reverse('marvelec_app:panel_trabajador'))
        self.assertEqual(ReporteAvance.objects.get().total_puntos, 3)


class CorreccionTests(BaseEscenario):

    def test_flujo_observar_y_corregir(self):
        reporte = self._avance(self.trab_a, ['DT-001', 'DT-002'], ['tendido'])
        self._login("sup_a")
        self.client.post(reverse('marvelec_app:revisar_reporte', args=[reporte.pk]),
                         {'accion': 'observar', 'comentario': 'DT-002 no está tendido'})
        reporte.refresh_from_db()
        self.assertEqual(reporte.estado, ReporteAvance.ESTADO_OBSERVADO)

        self._login("trab_a")
        r = self.client.get(reverse('marvelec_app:editar_reporte', args=[reporte.pk]))
        self.assertEqual(r.status_code, 200)
        r = self.client.post(reverse('marvelec_app:editar_reporte', args=[reporte.pk]),
                             {'rack': self.rack_a.pk, 'etapas': ['tendido'], 'puntos': [self._p('DT-001').pk]})
        self.assertRedirects(r, reverse('marvelec_app:reporte_detalle', args=[reporte.pk]))
        reporte.refresh_from_db()
        self.assertEqual(reporte.estado, ReporteAvance.ESTADO_PENDIENTE)
        self.assertEqual(list(reporte.avances.values_list('punto__numero', flat=True)), [1])
        self.assertTrue(Notificacion.objects.filter(tipo=Notificacion.TIPO_REPORTE_CORREGIDO, usuario=self.sup_a).exists())

    def test_no_se_quita_etapa_de_la_que_dependen_otros(self):
        reporte = self._avance(self.trab_a, ['DT-001', 'DT-002'], ['tendido'], estado=ReporteAvance.ESTADO_OBSERVADO)
        self._avance(self.trab_a2, ['DT-002'], ['bajada'])
        self._login("trab_a")
        self.client.post(reverse('marvelec_app:editar_reporte', args=[reporte.pk]),
                         {'rack': self.rack_a.pk, 'etapas': ['tendido'], 'puntos': [self._p('DT-001').pk]})
        self.assertEqual(reporte.avances.count(), 2)  # no se aplicó la corrección

    def test_solo_el_autor_corrige_y_solo_si_esta_observado(self):
        reporte = self._avance(self.trab_a, ['DT-001'], ['tendido'])
        self._login("trab_a")
        self.assertEqual(self.client.get(reverse('marvelec_app:editar_reporte', args=[reporte.pk])).status_code, 404)
        reporte.observar(self.sup_a, 'x')
        self._login("trab_a2")
        self.assertEqual(self.client.get(reverse('marvelec_app:editar_reporte', args=[reporte.pk])).status_code, 404)


# =========================================================
# Incidencias
# =========================================================

class IncidenciaTests(BaseEscenario):

    def setUp(self):
        super().setUp()
        self.sin_canalizacion = TipoIncidencia.objects.get(nombre='Sin canalización')

    def _post(self, **extra):
        datos = {'rack': self.rack_a.pk, 'codigos': 'DT 1-2', 'tipo': self.sin_canalizacion.pk,
                 'comentario': 'tubo tapado', 'latitud': '-33.45', 'longitud': '-70.66', **extra}
        return self.client.post(reverse('marvelec_app:api_incidencia'), datos)

    def test_reportar_bloquea_puntos_y_avisa(self):
        self._login("trab_a")
        r = self._post(foto=imagen())
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()['codigos'], ['DT-001', 'DT-002'])
        incidencias = list(Incidencia.objects.order_by('punto__numero'))
        self.assertEqual(incidencias[0].foto.name, incidencias[1].foto.name)  # una foto para ambos
        self.assertEqual(str(incidencias[0].latitud), '-33.450000')
        self.assertTrue(Notificacion.objects.filter(usuario=self.sup_a, tipo=Notificacion.TIPO_NUEVA_INCIDENCIA).exists())
        self.assertFalse(Notificacion.objects.filter(usuario=self.sup_b).exists())

    def test_no_duplica_incidencia_abierta_del_mismo_tipo(self):
        self._login("trab_a")
        self._post()
        r = self._post(codigos='DT 2-3')
        self.assertEqual(r.json()['codigos'], ['DT-003'])
        self.assertEqual(len(r.json()['omitidos']), 1)

    def test_liberar_desbloquea_y_mide_tiempo_detenido(self):
        self._login("trab_a")
        self._post()
        inc = Incidencia.objects.get(punto__numero=1)
        Incidencia.objects.filter(pk=inc.pk).update(creada=timezone.now() - timedelta(hours=5))

        self._login("sup_b")
        self.assertEqual(self.client.post(reverse('marvelec_app:liberar_incidencia', args=[inc.pk])).status_code, 404)
        self._login("sup_a")
        self.client.post(reverse('marvelec_app:liberar_incidencia', args=[inc.pk]), {'comentario': 'se despejó'})
        inc.refresh_from_db()
        self.assertEqual(inc.estado, Incidencia.ESTADO_LIBERADA)
        self.assertAlmostEqual(inc.horas_detenido, 5, delta=0.1)
        self.assertTrue(Notificacion.objects.filter(usuario=self.trab_a, tipo=Notificacion.TIPO_INCIDENCIA_LIBERADA).exists())

        self._login("trab_a")
        self.assertEqual(self._post_avance(codigos='DT 1').status_code, 201)

    def test_tipo_obligatorio(self):
        self._login("trab_a")
        self.assertEqual(self._post(tipo='').status_code, 400)


# =========================================================
# Asistencia
# =========================================================

class AsistenciaTests(BaseEscenario):

    def _marcar(self, **extra):
        datos = {'obra': self.obra_a.pk, 'foto': imagen(), 'latitud': '-33.4', 'longitud': '-70.6', 'precision': 12, **extra}
        return self.client.post(reverse('marvelec_app:api_asistencia'), datos)

    def test_marcar_llegada_una_vez_por_dia(self):
        self._login("trab_a")
        r = self._marcar()
        self.assertEqual(r.status_code, 201, r.content)
        a = Asistencia.objects.get()
        self.assertEqual((a.estado, a.fecha, a.precision), (Asistencia.ESTADO_PENDIENTE, consultas.hoy_local(), 12))
        self.assertEqual(self._marcar().status_code, 400)

    def test_foto_obligatoria(self):
        self._login("trab_a")
        r = self.client.post(reverse('marvelec_app:api_asistencia'), {'obra': self.obra_a.pk})
        self.assertEqual(r.status_code, 400)
        self.assertIn('foto', r.json()['mensaje'])

    def test_rechazo_avisa_y_permite_volver_a_marcar(self):
        self._login("trab_a")
        self._marcar()
        a = Asistencia.objects.get()
        self._login("sup_a")
        r = self.client.post(reverse('marvelec_app:revisar_asistencia', args=[a.pk]),
                             {'accion': 'rechazar', 'motivo': 'La foto no muestra la obra'},
                             HTTP_X_REQUESTED_WITH='fetch')
        self.assertEqual(r.json()['estado'], Asistencia.ESTADO_RECHAZADA)
        self.assertTrue(Notificacion.objects.filter(usuario=self.trab_a, tipo=Notificacion.TIPO_ASISTENCIA_RECHAZADA).exists())

        self._login("trab_a")
        self.assertEqual(self._marcar().status_code, 201)
        a.refresh_from_db()
        self.assertEqual(a.estado, Asistencia.ESTADO_PENDIENTE)

    def test_rechazo_exige_motivo_y_permisos_por_obra(self):
        self._login("trab_a")
        self._marcar()
        a = Asistencia.objects.get()
        self._login("sup_a")
        r = self.client.post(reverse('marvelec_app:revisar_asistencia', args=[a.pk]), {'accion': 'rechazar'},
                             HTTP_X_REQUESTED_WITH='fetch')
        self.assertEqual(r.status_code, 400)
        self._login("sup_b")
        r = self.client.post(reverse('marvelec_app:revisar_asistencia', args=[a.pk]), {'accion': 'confirmar'})
        self.assertEqual(r.status_code, 404)

    def test_confirmacion_masiva_y_vista_mensual(self):
        for u in ('trab_a', 'trab_a2'):
            self._login(u)
            self._marcar()
        self._login("sup_a")
        self.client.post(reverse('marvelec_app:confirmar_asistencias'),
                         {'ids': list(Asistencia.objects.values_list('pk', flat=True))})
        self.assertEqual(Asistencia.objects.filter(estado=Asistencia.ESTADO_CONFIRMADA).count(), 2)

        hoy = consultas.hoy_local()
        dias, filas = consultas.asistencia_mensual(PerfilTrabajador.objects.get(usuario=self.sup_a), hoy.year, hoy.month)
        self.assertEqual({f['perfil'].usuario.username: f['confirmados'] for f in filas}, {'trab_a': 1, 'trab_a2': 1})

        r = self.client.get(reverse('marvelec_app:exportar_asistencia'), {'mes': hoy.strftime('%Y-%m')})
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.cell(row=2, column=2 + len(dias) + 1).value, 1)


# =========================================================
# Permisos y revisión
# =========================================================

class PermisosTests(BaseEscenario):

    def test_supervisor_no_ve_reporte_ni_rack_de_otra_obra(self):
        reporte = self._avance(self.trab_b, ['DT-001'], ['tendido'], rack=self.rack_b)
        self._login("sup_a")
        self.assertEqual(self.client.get(reverse('marvelec_app:reporte_detalle', args=[reporte.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse('marvelec_app:rack_detalle', args=[self.rack_b.pk])).status_code, 404)
        punto = self._p('DT-001', self.rack_b)
        self.assertEqual(self.client.get(reverse('marvelec_app:punto_detalle', args=[punto.pk])).status_code, 404)

    def test_gerencia_ve_todo_y_trabajador_solo_lo_suyo(self):
        reporte = self._avance(self.trab_b, ['DT-001'], ['tendido'], rack=self.rack_b)
        self._login("gerente")
        self.assertEqual(self.client.get(reverse('marvelec_app:reporte_detalle', args=[reporte.pk])).status_code, 200)
        self._login("trab_a")
        self.assertEqual(self.client.get(reverse('marvelec_app:reporte_detalle', args=[reporte.pk])).status_code, 404)

    def test_trabajador_no_accede_a_supervision(self):
        self._login("trab_a")
        for nombre in ('panel_supervision', 'avance_obra', 'incidencias', 'asistencia_dia', 'cargar_puntos'):
            self.assertRedirects(self.client.get(reverse(f'marvelec_app:{nombre}')), reverse('marvelec_app:panel_trabajador'))

    def test_usuario_sin_perfil_recibe_uno_al_entrar(self):
        u = User.objects.create_user(username="sinperfil", password="clave12345")
        self.client.login(username="sinperfil", password="clave12345")
        self.client.get("/panel/", follow=True)
        self.assertTrue(PerfilTrabajador.objects.filter(usuario=u).exists())


class RevisionTests(BaseEscenario):

    def test_aprobar_notifica_al_trabajador(self):
        reporte = self._avance(self.trab_a, ['DT-001'], ['tendido'])
        self._login("sup_a")
        r = self.client.post(reverse('marvelec_app:revisar_reporte', args=[reporte.pk]), {'accion': 'aprobar'},
                             HTTP_X_REQUESTED_WITH='fetch')
        self.assertTrue(r.json()['ok'])
        reporte.refresh_from_db()
        self.assertEqual((reporte.estado, reporte.revisado_por), (ReporteAvance.ESTADO_APROBADO, self.sup_a))
        self.assertTrue(Notificacion.objects.filter(usuario=self.trab_a, tipo=Notificacion.TIPO_REPORTE_APROBADO).exists())

    def test_observar_exige_comentario(self):
        reporte = self._avance(self.trab_a, ['DT-001'], ['tendido'])
        self._login("sup_a")
        r = self.client.post(reverse('marvelec_app:revisar_reporte', args=[reporte.pk]), {'accion': 'observar'},
                             HTTP_X_REQUESTED_WITH='fetch')
        self.assertEqual(r.status_code, 400)

    def test_aprobacion_masiva_solo_de_reportes_visibles(self):
        propio = self._avance(self.trab_a, ['DT-001'], ['tendido'])
        ajeno = self._avance(self.trab_b, ['DT-001'], ['tendido'], rack=self.rack_b)
        self._login("sup_a")
        self.client.post(reverse('marvelec_app:revisar_masivo'), {'ids': [propio.pk, ajeno.pk]})
        propio.refresh_from_db()
        ajeno.refresh_from_db()
        self.assertEqual((propio.estado, ajeno.estado), (ReporteAvance.ESTADO_APROBADO, ReporteAvance.ESTADO_PENDIENTE))

    def test_cola_marca_reportes_sin_asistencia_confirmada(self):
        self._avance(self.trab_a, ['DT-001'], ['tendido'])
        self._login("sup_a")
        r = self.client.get(reverse('marvelec_app:cola_revision'))
        self.assertContains(r, 'Sin asistencia confirmada')


# =========================================================
# Indicadores, proyección y resumen diario
# =========================================================

class IndicadoresTests(BaseEscenario):

    def test_estado_por_rack_y_proyeccion(self):
        self._avance(self.trab_a, ['DT-001', 'DT-002', 'DT-003', 'DT-004', 'DT-005', 'DT-006'], ['tendido'])
        self._avance(self.trab_a, ['DT-001'], ['bajada'])
        Incidencia.objects.create(punto=self._p('VD-001'), tipo=TipoIncidencia.objects.first(), reportada_por=self.trab_a)

        datos = consultas.resumen_obra(self.obra_a)
        fila = datos['filas'][0]
        self.assertEqual((fila['total'], fila['bloqueados']), (7, 1))
        tendido = next(c for c in fila['etapas'] if c['clave'] == 'tendido')
        self.assertEqual((tendido['n'], tendido['pct']), (6, 86))

        proy = {p['etapa'].clave: p for p in datos['proyeccion']}
        self.assertEqual(proy['tendido']['faltan'], 1)
        self.assertGreater(proy['tendido']['ritmo'], 0)
        self.assertIsNotNone(proy['tendido']['fin'])
        self.assertIsNone(proy['certificacion']['fin'])
        self.assertEqual(proy['certificacion']['dias_persona'], round(7 / etapas.POR_CLAVE['certificacion'].estandar, 1))
        self.assertIsNone(datos['fin_estimado'])

    def test_rendimiento_contra_estandar_por_jornada_confirmada(self):
        hoy = consultas.hoy_local()
        self._avance(self.trab_a, ['DT-001', 'DT-002', 'DT-003', 'DT-004', 'DT-005'], ['bajada'])
        Asistencia.objects.create(trabajador=self.trab_a, obra=self.obra_a, fecha=hoy, foto='x.jpg',
                                  estado=Asistencia.ESTADO_CONFIRMADA)
        perfil = PerfilTrabajador.objects.get(usuario=self.sup_a)
        ranking = consultas.ranking_trabajadores(consultas.reportes_visibles(perfil), hoy, hoy)
        r = ranking[0]
        self.assertEqual((r['puntos'], r['jornadas'], r['jornadas_asistencia']), (5, 1, 1))
        self.assertEqual(r['rendimiento'], round(100 * 5 / etapas.POR_CLAVE['bajada'].estandar))

    def test_datos_del_selector(self):
        self._avance(self.trab_a, ['DT-001'], ['tendido', 'bajada'])
        Incidencia.objects.create(punto=self._p('DT-002'), tipo=TipoIncidencia.objects.first(), reportada_por=self.trab_a)
        datos = consultas.datos_selector(Obra.objects.filter(pk=self.obra_a.pk))
        rack = datos['racks'][0]
        puntos = {(p[1], p[2]): p for p in rack['puntos']}
        self.assertEqual(puntos[('DT', 1)][4], 0b11)  # tendido y bajada
        self.assertEqual(puntos[('DT', 2)][5], 1)  # bloqueado
        self.assertEqual(rack['paquetes'][1], [36, 'Sótano 1'])


class ResumenDiarioTests(BaseEscenario):

    def test_resumen_lista_quien_no_marco_y_no_se_duplica(self):
        from .resumen import enviar_resumen_diario, resumen_para
        hoy = consultas.hoy_local()
        self._avance(self.trab_a, ['DT-001', 'DT-002'], ['tendido'])
        Asistencia.objects.create(trabajador=self.trab_a, obra=self.obra_a, fecha=hoy, foto='x.jpg')
        datos = resumen_para(PerfilTrabajador.objects.get(usuario=self.sup_a), hoy)
        self.assertEqual(datos['puntos'], 2)
        self.assertEqual([p.usuario.username for p in datos['sin_marcar']], ['trab_a2'])
        self.assertEqual(datos['sin_confirmar'], 1)

        self.assertTrue(enviar_resumen_diario(hoy)['enviado'])
        self.assertFalse(enviar_resumen_diario(hoy)['enviado'])
        self.assertTrue(ResumenDiarioEnviado.objects.filter(fecha=hoy).exists())
        aviso = Notificacion.objects.get(usuario=self.sup_a, tipo=Notificacion.TIPO_RESUMEN_DIARIO)
        self.assertIn('Juan Pérez', aviso.cuerpo)

    @override_settings(CRON_TOKEN='secreto')
    def test_endpoint_requiere_token(self):
        url = reverse('marvelec_app:tarea_resumen_diario')
        self.assertEqual(self.client.post(url).status_code, 403)
        self.assertEqual(self.client.post(url, HTTP_X_CRON_TOKEN='secreto').status_code, 200)

    @override_settings(CRON_TOKEN='')
    def test_endpoint_deshabilitado_sin_token_configurado(self):
        self.assertEqual(self.client.post(reverse('marvelec_app:tarea_resumen_diario')).status_code, 403)


class NotificacionesTests(BaseEscenario):

    def test_estado_y_marcar_leida(self):
        n = Notificacion.objects.create(usuario=self.trab_a, tipo='reporte_aprobado', titulo='Aprobado', url='/panel/')
        self._login("trab_a")
        data = self.client.get(reverse('marvelec_app:notificaciones_estado')).json()
        self.assertEqual(data['no_leidas'], 1)
        r = self.client.post(reverse('marvelec_app:notificacion_leer', args=[n.pk]), HTTP_X_REQUESTED_WITH='fetch')
        self.assertEqual(r.json()['url'], '/panel/')
        n.refresh_from_db()
        self.assertTrue(n.leida)

    def test_no_se_puede_leer_notificacion_ajena(self):
        n = Notificacion.objects.create(usuario=self.trab_a2, tipo='reporte_aprobado', titulo='x')
        self._login("trab_a")
        self.assertEqual(self.client.post(reverse('marvelec_app:notificacion_leer', args=[n.pk])).status_code, 404)

    def test_push_se_envia_a_dispositivos_suscritos(self):
        SuscripcionPush.objects.create(usuario=self.sup_a, endpoint='https://push.example/1', p256dh='k', auth='a')
        with mock.patch('marvelec_app.push.enviar_notificacion_push') as enviar:
            self._login("trab_a")
            self._post_avance(codigos='DT 1')
        self.assertTrue(any(c.args[0] == self.sup_a for c in enviar.call_args_list))


# =========================================================
# Excel y carga de puntos desde la pantalla
# =========================================================

class ExcelTests(BaseEscenario):

    def test_excel_de_avance_una_fila_por_punto_y_etapa(self):
        self._avance(self.trab_a, ['DT-001', 'DT-002'], ['tendido', 'bajada'])
        self._avance(self.trab_a2, ['DT-003'], ['tendido'])
        self._login("sup_a")
        r = self.client.get(reverse('marvelec_app:exportar_excel'), {'trabajador': self.trab_a2.pk})
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws.cell(row=2, column=8).value, 'DT-003')

    def test_estado_de_puntos(self):
        self._avance(self.trab_a, ['DT-001'], ['tendido'])
        self._login("sup_a")
        r = self.client.get(reverse('marvelec_app:exportar_estado_puntos'), {'obra': self.obra_a.pk})
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.max_row, 1 + 7)  # encabezado + puntos del rack 1H
        fila = [c.value for c in ws[2]]
        self.assertEqual(fila[4], 'DT-001')
        self.assertEqual(fila[7], 'Pedro Soto')  # quién hizo el tendido


class CargarPuntosVistaTests(BaseEscenario):

    def _archivo(self):
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.title = 'RACK_3C'
        hoja['A1'], hoja['A2'], hoja['A3'], hoja['A4'] = 'PAQUETE 01', 'DT-001', 'DT-002', '45 METROS'
        buffer = io.BytesIO()
        libro.save(buffer)
        return SimpleUploadedFile('planilla.xlsx', buffer.getvalue())

    def test_vista_previa_y_confirmar(self):
        self._login("sup_a")
        r = self.client.post(reverse('marvelec_app:cargar_puntos'), {'obra': self.obra_a.pk, 'archivo': self._archivo()})
        self.assertContains(r, 'Vista previa')
        self.assertEqual(r.context['comparacion']['nuevos'], 2)
        self.assertFalse(Rack.objects.filter(nombre='3C').exists())

        r = self.client.post(reverse('marvelec_app:cargar_puntos'), {'confirmar': '1', 'datos': r.context['datos']})
        self.assertRedirects(r, reverse('marvelec_app:avance_obra') + f'?obra={self.obra_a.pk}')
        self.assertEqual(Punto.objects.filter(rack__nombre='3C', rack__obra=self.obra_a).count(), 2)

    def test_datos_alterados_o_de_otra_obra_se_rechazan(self):
        self._login("sup_a")
        r = self.client.post(reverse('marvelec_app:cargar_puntos'), {'confirmar': '1', 'datos': 'basura'})
        self.assertRedirects(r, reverse('marvelec_app:cargar_puntos'))
        datos = signing.dumps({'obra': self.obra_b.pk, 'puntos': []}, salt='marvelec.cargar-puntos', compress=True)
        self.client.post(reverse('marvelec_app:cargar_puntos'), {'confirmar': '1', 'datos': datos})
        self.assertFalse(Rack.objects.filter(obra=self.obra_b).exclude(pk=self.rack_b.pk).exists())

    def test_solo_xlsx(self):
        self._login("sup_a")
        r = self.client.post(reverse('marvelec_app:cargar_puntos'),
                             {'obra': self.obra_a.pk, 'archivo': SimpleUploadedFile('a.csv', b'x')})
        self.assertContains(r, '.xlsx')


# =========================================================
# Todas las pantallas cargan
# =========================================================

class PantallasTests(BaseEscenario):

    def setUp(self):
        super().setUp()
        self.reporte = self._avance(self.trab_a, ['DT-001', 'DT-002'], ['tendido'], comentario='Todo bien')
        self._avance(self.trab_b, ['DT-001'], ['tendido'], rack=self.rack_b)
        self.incidencia = Incidencia.objects.create(punto=self._p('DT-003'), tipo=TipoIncidencia.objects.first(),
                                                    reportada_por=self.trab_a)
        Asistencia.objects.create(trabajador=self.trab_a, obra=self.obra_a, fecha=consultas.hoy_local(), foto='x.jpg')
        Notificacion.objects.create(usuario=self.trab_a, tipo='reporte_aprobado', titulo='Aprobado', cuerpo='ok')

    def _ok(self, nombre, *args, params=None):
        r = self.client.get(reverse(f'marvelec_app:{nombre}', args=args), params or {})
        self.assertEqual(r.status_code, 200, f'{nombre} {params or ""} -> {r.status_code}')
        return r

    def test_pantallas_trabajador(self):
        self._login("trab_a")
        for nombre in ('panel_trabajador', 'marcar_asistencia', 'nuevo_reporte', 'nueva_incidencia',
                       'historial_trabajador', 'notificaciones'):
            self._ok(nombre)
        for ver in ('llegadas', 'incidencias'):
            self._ok('historial_trabajador', params={'ver': ver})
        self._ok('historial_trabajador', params={'estado': 'pendiente'})
        r = self._ok('reporte_detalle', self.reporte.pk)
        self.assertContains(r, 'DT-002')
        r = self._ok('nuevo_reporte')
        self.assertContains(r, 'selector-data')

    def test_pantallas_supervision(self):
        for usuario in ('sup_a', 'gerente'):
            self._login(usuario)
            r = self._ok('panel_supervision')
            self.assertContains(r, 'datos-graficos')
            for rango in ('hoy', '7', 'mes'):
                self._ok('panel_supervision', params={'rango': rango})
            self._ok('panel_supervision', params={'desde': '2020-01-01', 'hasta': '2030-01-01', 'obra': self.obra_a.pk,
                                                  'etapa': 'tendido', 'rack': self.rack_a.pk})
            r = self._ok('avance_obra', params={'obra': self.obra_a.pk})
            self.assertContains(r, 'Proyección por etapa')
            for params in ({}, {'pendiente': 'bajada'}, {'bloqueados': '1'}):
                self._ok('rack_detalle', self.rack_a.pk, params=params)
            self._ok('punto_detalle', self._p('DT-003').pk)
            for estado in ('abierta', 'liberada'):
                self._ok('incidencias', params={'estado': estado})
            self._ok('asistencia_dia')
            self._ok('asistencia_mes')
            self._ok('asistencia_mes', params={'mes': '2026-02'})
            for nombre in ('cola_revision', 'lista_reportes', 'trabajadores', 'notificaciones', 'cargar_puntos'):
                self._ok(nombre)
            self._ok('lista_reportes', params={'estado': 'pendiente', 'obra': self.obra_a.pk, 'etapa': 'tendido'})
            r = self._ok('reporte_detalle', self.reporte.pk)
            self.assertContains(r, 'Aprobar')
        self.assertEqual(self.client.get(reverse('marvelec_app:plantilla_puntos')).status_code, 200)

    def test_avance_obra_sin_puntos_invita_a_cargar(self):
        Obra.objects.create(nombre='Obra vacía')
        self._login("gerente")
        r = self._ok('avance_obra', params={'obra': Obra.objects.get(nombre='Obra vacía').pk})
        self.assertContains(r, 'Cargar puntos')

    def test_pwa_y_login(self):
        r = self._ok('service_worker')
        self.assertEqual(r['Content-Type'], 'application/javascript; charset=utf-8')
        self.assertContains(r, 'importScripts')
        self.assertEqual(json.loads(self._ok('manifest').content)['short_name'], 'MARVELEC')
        self._ok('login')

    def test_logout_por_post(self):
        self._login("trab_a")
        r = self.client.post(reverse('marvelec_app:logout'))
        self.assertRedirects(r, reverse('marvelec_app:login') + '?salida=1')
        self.assertRedirects(self.client.get(reverse('marvelec_app:panel')), reverse('marvelec_app:login') + '?next=/panel/')


# =========================================================
# Panel de administración (tema Unfold)
# =========================================================

class AdminTests(BaseEscenario):

    def setUp(self):
        super().setUp()
        self.admin = User.objects.create_superuser('admin', 'admin@marvelec.cl', 'clave12345')

    def test_listados_del_admin_cargan(self):
        reporte = self._avance(self.trab_a, ['DT-001'], ['tendido'])
        Incidencia.objects.create(punto=self._p('DT-002'), tipo=TipoIncidencia.objects.first(), reportada_por=self.trab_a)
        Asistencia.objects.create(trabajador=self.trab_a, obra=self.obra_a, fecha=consultas.hoy_local(), foto='x.jpg')
        self._login('admin')
        for nombre in ('admin:index', 'admin:auth_user_changelist', 'admin:auth_group_changelist'):
            self.assertEqual(self.client.get(reverse(nombre)).status_code, 200, nombre)
        for modelo in ('obra', 'zona', 'rack', 'paquete', 'punto', 'tipopunto', 'perfiltrabajador', 'reporteavance',
                       'tipoincidencia', 'incidencia', 'asistencia', 'notificacion', 'suscripcionpush',
                       'resumendiarioenviado'):
            r = self.client.get(reverse(f'admin:marvelec_app_{modelo}_changelist'))
            self.assertEqual(r.status_code, 200, modelo)
        for modelo, pk in (('reporteavance', reporte.pk), ('punto', self._p('DT-001').pk), ('rack', self.rack_a.pk)):
            self.assertEqual(self.client.get(reverse(f'admin:marvelec_app_{modelo}_change', args=[pk])).status_code, 200)

        r = self.client.get(reverse('admin:marvelec_app_rack_changelist'))
        totales = {x.nombre: (x._n_puntos, x._n_cert) for x in r.context['cl'].result_list}
        self.assertEqual(totales, {'1H': (7, 0), '2A': (3, 0)})
        r = self.client.get(reverse('admin:marvelec_app_punto_changelist'), {'q': 'dt-1'})
        self.assertEqual(r.context['cl'].result_count, 2)  # DT-001 de cada rack

    def test_acciones_aprobar_liberar_y_confirmar(self):
        pendiente = self._avance(self.trab_a, ['DT-001'], ['tendido'])
        inc = Incidencia.objects.create(punto=self._p('DT-002'), tipo=TipoIncidencia.objects.first(), reportada_por=self.trab_a)
        asis = Asistencia.objects.create(trabajador=self.trab_a, obra=self.obra_a, fecha=consultas.hoy_local(), foto='x.jpg')
        self._login('admin')
        self.client.post(reverse('admin:marvelec_app_reporteavance_changelist'),
                         {'action': 'aprobar_seleccionados', '_selected_action': [pendiente.pk]})
        self.client.post(reverse('admin:marvelec_app_incidencia_changelist'),
                         {'action': 'liberar_seleccionadas', '_selected_action': [inc.pk]})
        self.client.post(reverse('admin:marvelec_app_asistencia_changelist'),
                         {'action': 'confirmar_seleccionadas', '_selected_action': [asis.pk]})
        pendiente.refresh_from_db()
        inc.refresh_from_db()
        asis.refresh_from_db()
        self.assertEqual(pendiente.estado, ReporteAvance.ESTADO_APROBADO)
        self.assertEqual(inc.estado, Incidencia.ESTADO_LIBERADA)
        self.assertEqual(asis.estado, Asistencia.ESTADO_CONFIRMADA)
