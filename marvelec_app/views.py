# marvelec_app/views.py
from datetime import timedelta
from functools import wraps
import hmac

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db.models import Max, Sum
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.utils.timesince import timesince
from django.views.decorators.cache import never_cache
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST

from . import consultas
from .forms import ReporteAvanceForm, PuntosForm, RevisionForm
from .models import Notificacion, PerfilTrabajador, ReporteAvance, RegistroPunto
from .servicios import crear_reporte, corregir_reporte, revisar_reporte, leer_uuid

# Cambiar al modificar sw.js para forzar su actualización en los dispositivos.
SW_VERSION = '2.0.0'


# =========================================================
# AUTENTICACIÓN
# =========================================================

def login_usuario(request):
    if request.user.is_authenticated:
        return redirect('marvelec_app:panel')

    siguiente = request.POST.get('next') or request.GET.get('next') or ''
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')
        user = authenticate(request, username=username, password=password)
        if user is not None:
            login(request, user)
            if siguiente and url_has_allowed_host_and_scheme(siguiente, {request.get_host()}):
                return redirect(siguiente)
            return redirect('marvelec_app:panel')
        messages.error(request, 'Usuario o contraseña incorrectos.')

    return render(request, 'registration/login.html', {'next': siguiente})


def logout_usuario(request):
    logout(request)
    return redirect(reverse('marvelec_app:login') + '?salida=1')


def _get_perfil(user):
    """Devuelve el PerfilTrabajador del usuario, creándolo si por algún motivo no existe."""
    perfil, _ = PerfilTrabajador.objects.select_related('obra_asignada').get_or_create(usuario=user)
    return perfil


def revisor_requerido(vista):
    """Sólo supervisores y gerencia; un trabajador vuelve a su app."""
    @wraps(vista)
    @login_required
    def envoltura(request, *args, **kwargs):
        perfil = _get_perfil(request.user)
        if not perfil.es_supervisor_o_gerencia:
            return redirect('marvelec_app:panel_trabajador')
        request.perfil = perfil
        return vista(request, *args, **kwargs)
    return envoltura


def _es_fetch(request):
    return request.headers.get('X-Requested-With') == 'fetch'


@login_required
def dashboard_redirect(request):
    perfil = _get_perfil(request.user)
    if perfil.rol == PerfilTrabajador.ROL_TRABAJADOR:
        return redirect('marvelec_app:panel_trabajador')
    return redirect('marvelec_app:panel_supervision')


# =========================================================
# APP DEL TRABAJADOR
# =========================================================

@login_required
def dashboard_trabajador(request):
    perfil = _get_perfil(request.user)
    hoy = consultas.hoy_local()
    inicio_semana = hoy - timedelta(days=hoy.weekday())
    mis_reportes = ReporteAvance.objects.filter(trabajador=request.user)

    context = {
        'perfil': perfil,
        'hoy': hoy,
        'puntos_hoy': mis_reportes.filter(fecha_hora__date=hoy).aggregate(t=Sum('total_puntos'))['t'] or 0,
        'puntos_semana': mis_reportes.filter(fecha_hora__date__gte=inicio_semana).aggregate(t=Sum('total_puntos'))['t'] or 0,
        'reportes_semana': mis_reportes.filter(fecha_hora__date__gte=inicio_semana).count(),
        'reporto_hoy': mis_reportes.filter(fecha_hora__date=hoy).exists(),
        'observados': mis_reportes.filter(estado=ReporteAvance.ESTADO_OBSERVADO).select_related('obra', 'subetapa'),
        'recientes': mis_reportes.select_related('obra', 'subetapa').prefetch_related('registros')[:5],
    }
    return render(request, 'marvelec_app/trabajador/inicio.html', context)


def _contexto_formulario(perfil, form, puntos_form):
    obras = form.fields['obra'].queryset
    return {
        'perfil': perfil,
        'form': form,
        'puntos_form': puntos_form,
        'subetapas_por_obra': consultas.subetapas_por_obra(obras),
        'iconos_tipo': {'red': 'network', 'fuerza': 'plug', 'iluminacion': 'bulb'},
    }


@login_required
def nuevo_reporte(request):
    """Formulario de reporte. Con JS se envía por la API (con soporte sin señal); sin JS, por aquí."""
    perfil = _get_perfil(request.user)

    if request.method == 'POST':
        form = ReporteAvanceForm(request.POST, request.FILES, obra_asignada=perfil.obra_asignada)
        puntos_form = PuntosForm(request.POST)
        if form.is_valid() and puntos_form.is_valid():
            try:
                crear_reporte(request.user, form, puntos_form)
                messages.success(request, 'Reporte de avance enviado correctamente.')
                return redirect('marvelec_app:panel_trabajador')
            except ValidationError as e:
                for msg in e.messages:
                    messages.error(request, msg)
        else:
            messages.error(request, 'Revisa los datos del reporte: hay errores en el formulario.')
    else:
        form = ReporteAvanceForm(obra_asignada=perfil.obra_asignada)
        puntos_form = PuntosForm()

    return render(request, 'marvelec_app/trabajador/nuevo_reporte.html',
                  _contexto_formulario(perfil, form, puntos_form))


@login_required
def historial_trabajador(request):
    reportes = ReporteAvance.objects.filter(trabajador=request.user).select_related(
        'obra', 'subetapa'
    ).prefetch_related('registros')
    estado = request.GET.get('estado', '')
    if estado in dict(ReporteAvance.ESTADO_CHOICES):
        reportes = reportes.filter(estado=estado)
    pagina = Paginator(reportes, 20).get_page(request.GET.get('page'))
    return render(request, 'marvelec_app/trabajador/historial.html', {
        'pagina': pagina,
        'estado': estado,
        'estados': ReporteAvance.ESTADO_CHOICES,
    })


@login_required
def editar_reporte(request, pk):
    """El trabajador corrige un reporte que el supervisor observó."""
    reporte = get_object_or_404(ReporteAvance.objects.select_related('obra', 'subetapa'), pk=pk)
    if not reporte.puede_editar(request.user):
        raise Http404
    perfil = _get_perfil(request.user)

    if request.method == 'POST':
        form = ReporteAvanceForm(request.POST, request.FILES, instance=reporte, obra_asignada=perfil.obra_asignada)
        puntos_form = PuntosForm(request.POST, reporte=reporte)
        if form.is_valid() and puntos_form.is_valid():
            try:
                corregir_reporte(reporte, form, puntos_form)
                messages.success(request, 'Reporte corregido y enviado nuevamente a revisión.')
                return redirect('marvelec_app:reporte_detalle', pk=reporte.pk)
            except ValidationError as e:
                for msg in e.messages:
                    messages.error(request, msg)
        else:
            messages.error(request, 'Revisa los datos del reporte.')
    else:
        form = ReporteAvanceForm(instance=reporte, obra_asignada=perfil.obra_asignada)
        puntos_form = PuntosForm(reporte=reporte)

    context = _contexto_formulario(perfil, form, puntos_form)
    context['reporte'] = reporte
    return render(request, 'marvelec_app/trabajador/nuevo_reporte.html', context)


@require_POST
def api_crear_reporte(request):
    """
    Recibe un reporte desde la app del celular (también los que se guardaron sin señal).
    Responde JSON. Es idempotente gracias a `uuid_cliente`.
    """
    if not request.user.is_authenticated:
        return JsonResponse({'ok': False, 'mensaje': 'Sesión expirada.'}, status=401)

    # Un reporte guardado en el teléfono por otro usuario no se envía a nombre de quien ingresó ahora.
    usuario_id = request.POST.get('usuario_id')
    if usuario_id and usuario_id != str(request.user.id):
        return JsonResponse({'ok': False, 'mensaje': 'Este reporte pertenece a otro usuario.'}, status=409)

    perfil = _get_perfil(request.user)
    form = ReporteAvanceForm(request.POST, request.FILES, obra_asignada=perfil.obra_asignada)
    puntos_form = PuntosForm(request.POST)

    if not (form.is_valid() and puntos_form.is_valid()):
        errores = {**form.errors.get_json_data(), **puntos_form.errors.get_json_data()}
        primero = next((e[0]['message'] for e in errores.values() if e), 'Datos inválidos.')
        return JsonResponse({'ok': False, 'mensaje': primero, 'errores': errores}, status=400)

    try:
        reporte, creado = crear_reporte(
            request.user, form, puntos_form,
            uuid_cliente=leer_uuid(request.POST.get('uuid_cliente')),
            fecha_cliente=request.POST.get('fecha_hora_cliente'),
            offline=request.POST.get('offline') == '1',
        )
    except ValidationError as e:
        return JsonResponse({'ok': False, 'mensaje': ' '.join(e.messages)}, status=400)

    return JsonResponse({
        'ok': True,
        'id': reporte.pk,
        'duplicado': not creado,
        'total_puntos': reporte.total_puntos,
        'url': reverse('marvelec_app:reporte_detalle', args=[reporte.pk]),
    }, status=201 if creado else 200)


# =========================================================
# DETALLE DE REPORTE (trabajador dueño, supervisor de la obra, gerencia)
# =========================================================

@login_required
def reporte_detalle(request, pk):
    perfil = _get_perfil(request.user)
    reporte = get_object_or_404(
        ReporteAvance.objects.select_related('trabajador', 'obra', 'subetapa', 'revisado_por')
        .prefetch_related('registros'),
        pk=pk,
    )
    if not reporte.puede_ver(perfil):
        raise Http404

    puede_revisar = perfil.puede_revisar_obra(reporte.obra_id) and reporte.trabajador_id != request.user.id
    context = {
        'reporte': reporte,
        'perfil': perfil,
        'puntos': reporte.puntos_por_tipo(),
        'puede_revisar': puede_revisar,
        'puede_editar': reporte.puede_editar(request.user),
        'revision_form': RevisionForm() if puede_revisar else None,
        'siguiente_pendiente': None,
    }
    if puede_revisar:
        context['siguiente_pendiente'] = consultas.reportes_visibles(perfil).filter(
            estado=ReporteAvance.ESTADO_PENDIENTE
        ).exclude(pk=reporte.pk).order_by('fecha_hora').first()

    plantilla = 'marvelec_app/reporte_detalle.html' if perfil.es_supervisor_o_gerencia \
        else 'marvelec_app/trabajador/reporte_detalle.html'
    return render(request, plantilla, context)


@require_POST
@revisor_requerido
def revisar(request, pk):
    """Aprobar u observar un reporte (formulario normal o fetch desde la cola de revisión)."""
    reporte = get_object_or_404(ReporteAvance.objects.select_related('trabajador', 'obra', 'subetapa'), pk=pk)
    if not request.perfil.puede_revisar_obra(reporte.obra_id) or reporte.trabajador_id == request.user.id:
        raise Http404

    form = RevisionForm(request.POST)
    if not form.is_valid():
        mensaje = ' '.join(form.non_field_errors()) or 'Acción inválida.'
        if _es_fetch(request):
            return JsonResponse({'ok': False, 'mensaje': mensaje}, status=400)
        messages.error(request, mensaje)
        return redirect('marvelec_app:reporte_detalle', pk=pk)

    revisar_reporte(reporte, request.user, form.cleaned_data['accion'], form.cleaned_data['comentario'])
    texto = 'Reporte aprobado.' if reporte.estado == ReporteAvance.ESTADO_APROBADO else 'Reporte observado; se avisó al trabajador.'

    if _es_fetch(request):
        return JsonResponse({'ok': True, 'estado': reporte.estado, 'mensaje': texto})
    messages.success(request, texto)
    siguiente = request.POST.get('next')
    if siguiente and url_has_allowed_host_and_scheme(siguiente, {request.get_host()}):
        return redirect(siguiente)
    return redirect('marvelec_app:reporte_detalle', pk=pk)


@require_POST
@revisor_requerido
def revisar_masivo(request):
    """Aprueba varios reportes pendientes a la vez."""
    ids = [i for i in request.POST.getlist('ids') if i.isdigit()]
    reportes = consultas.reportes_visibles(request.perfil).filter(
        pk__in=ids, estado=ReporteAvance.ESTADO_PENDIENTE
    ).exclude(trabajador=request.user)
    n = 0
    for reporte in reportes:
        revisar_reporte(reporte, request.user, 'aprobar')
        n += 1
    texto = f'{n} reporte(s) aprobados.'
    if _es_fetch(request):
        return JsonResponse({'ok': True, 'aprobados': n, 'mensaje': texto})
    messages.success(request, texto)
    return redirect('marvelec_app:cola_revision')


# =========================================================
# PROGRAMA DE SUPERVISIÓN / GERENCIA
# =========================================================

def _contexto_filtros(request, perfil, filtros):
    obras = consultas.obras_visibles(perfil)
    obra_fija = perfil.rol == PerfilTrabajador.ROL_SUPERVISOR and perfil.obra_asignada_id
    return {
        'filtros': filtros,
        'obras_disponibles': obras,
        'obra_fija': obra_fija,
        'subetapas_por_obra': consultas.subetapas_por_obra(obras),
        'trabajadores_disponibles': consultas.trabajadores_de(perfil),
        'estados': ReporteAvance.ESTADO_CHOICES,
    }


def _querystring(filtros):
    from urllib.parse import urlencode
    params = {k: (v.isoformat() if hasattr(v, 'isoformat') else v) for k, v in filtros.items() if v}
    return urlencode(params)


@revisor_requerido
def dashboard_supervisor(request):
    perfil = request.perfil
    hoy = consultas.hoy_local()
    # Atajos de período (Hoy / 7 días / 30 días / Este mes); por defecto, últimos 30 días.
    rangos = {'hoy': hoy, '7': hoy - timedelta(days=6), '30': hoy - timedelta(days=29), 'mes': hoy.replace(day=1)}
    if 'desde' in request.GET:
        rango = ''  # fechas elegidas a mano
    else:
        rango = request.GET.get('rango', '30')
        rango = rango if rango in rangos else '30'
    filtros = consultas.leer_filtros(request.GET, desde_por_defecto=rangos[rango] if rango else None)
    hasta = filtros['hasta'] or hoy

    base = consultas.reportes_visibles(perfil)
    reportes = consultas.aplicar_filtros(base, filtros)

    total_puntos = reportes.aggregate(t=Sum('total_puntos'))['t'] or 0
    ranking = consultas.ranking_trabajadores(reportes)
    jornadas = sum(r['jornadas'] for r in ranking)

    # Hoy: quién reportó y quién no (según los trabajadores a cargo).
    a_cargo = consultas.trabajadores_de(perfil)
    if filtros['obra']:
        a_cargo = a_cargo.filter(obra_asignada_id=filtros['obra'])
    reportaron_hoy = set(base.filter(fecha_hora__date=hoy).values_list('trabajador_id', flat=True))
    sin_reportar = [p for p in a_cargo if p.usuario_id not in reportaron_hoy]

    tipos = dict(RegistroPunto.TIPO_CHOICES)
    por_tipo = consultas.puntos_por_tipo(reportes)
    context = {
        **_contexto_filtros(request, perfil, filtros),
        'perfil': perfil,
        'hoy': hoy,
        'total_puntos': total_puntos,
        'total_reportes': reportes.count(),
        'pendientes': base.filter(estado=ReporteAvance.ESTADO_PENDIENTE).count(),
        'promedio_jornada': round(total_puntos / jornadas, 1) if jornadas else 0,
        'reportaron_hoy': len(reportaron_hoy),
        'a_cargo': len(a_cargo),
        'sin_reportar': sin_reportar,
        'top': ranking[:5],
        'max_top': max((r['puntos_jornada'] for r in ranking[:5]), default=0),
        'recientes': reportes.order_by('-fecha_hora')[:8],
        'querystring': _querystring(filtros),
        'graficos': {
            'por_dia': consultas.serie_puntos_por_dia(reportes, filtros['desde'], hasta),
            'por_subetapa': consultas.puntos_por_subetapa(reportes),
            'por_tipo': {'labels': [tipos[t] for t in por_tipo], 'valores': list(por_tipo.values()), 'claves': list(por_tipo)},
        },
        'rango': rango,
    }
    return render(request, 'marvelec_app/supervision/resumen.html', context)


@revisor_requerido
def cola_revision(request):
    perfil = request.perfil
    filtros = consultas.leer_filtros(request.GET)
    pendientes = consultas.aplicar_filtros(consultas.reportes_visibles(perfil), filtros).filter(
        estado=ReporteAvance.ESTADO_PENDIENTE
    ).exclude(trabajador=request.user).order_by('fecha_hora')
    pagina = Paginator(pendientes, 20).get_page(request.GET.get('page'))
    return render(request, 'marvelec_app/supervision/revisar.html', {
        **_contexto_filtros(request, perfil, filtros),
        'pagina': pagina,
        'total': pagina.paginator.count,
    })


@revisor_requerido
def lista_reportes(request):
    perfil = request.perfil
    filtros = consultas.leer_filtros(request.GET)
    reportes = consultas.aplicar_filtros(consultas.reportes_visibles(perfil), filtros).order_by('-fecha_hora')
    pagina = Paginator(reportes, 25).get_page(request.GET.get('page'))
    return render(request, 'marvelec_app/supervision/reportes.html', {
        **_contexto_filtros(request, perfil, filtros),
        'pagina': pagina,
        'total_puntos': reportes.aggregate(t=Sum('total_puntos'))['t'] or 0,
        'querystring': _querystring(filtros),
    })


@revisor_requerido
def trabajadores(request):
    perfil = request.perfil
    hoy = consultas.hoy_local()
    filtros = consultas.leer_filtros(request.GET, desde_por_defecto=hoy.replace(day=1))
    reportes = consultas.aplicar_filtros(consultas.reportes_visibles(perfil), filtros)
    ranking = consultas.ranking_trabajadores(reportes)

    # Incluir a los trabajadores a cargo que no tienen reportes en el período.
    con_reportes = {r['id'] for r in ranking}
    a_cargo = consultas.trabajadores_de(perfil)
    if filtros['obra']:
        a_cargo = a_cargo.filter(obra_asignada_id=filtros['obra'])
    sin_actividad = [p for p in a_cargo if p.usuario_id not in con_reportes]

    ultimos = {}
    for fila in reportes.order_by().values('trabajador_id').annotate(ultimo=Max('fecha_hora')):
        ultimos[fila['trabajador_id']] = fila['ultimo']
    for r in ranking:
        r['ultimo'] = ultimos.get(r['id'])

    return render(request, 'marvelec_app/supervision/trabajadores.html', {
        **_contexto_filtros(request, perfil, filtros),
        'ranking': ranking,
        'max_puntos_jornada': max((r['puntos_jornada'] for r in ranking), default=0),
        'sin_actividad': sin_actividad,
        'querystring': _querystring(filtros),
    })


@revisor_requerido
def exportar_reportes_excel(request):
    """Exporta a Excel los reportes visibles para el usuario (con los mismos filtros de la pantalla)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    filtros = consultas.leer_filtros(request.GET)
    reportes = consultas.aplicar_filtros(consultas.reportes_visibles(request.perfil), filtros).order_by('-fecha_hora')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reportes de Avance"

    headers = ['Fecha', 'Trabajador', 'Obra', 'Subetapa', 'Puntos Red', 'Puntos Fuerza',
               'Puntos Iluminación', 'Total Puntos', 'Estado', 'Revisado por', 'Comentario', 'Observación']
    ws.append(headers)
    relleno = PatternFill('solid', fgColor='3A288F')
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = relleno
        celda.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 22

    totales = [0, 0, 0, 0]
    for reporte in reportes:
        conteo = reporte.puntos_por_tipo()
        fila = [
            conteo.get(RegistroPunto.TIPO_RED, 0),
            conteo.get(RegistroPunto.TIPO_FUERZA, 0),
            conteo.get(RegistroPunto.TIPO_ILUMINACION, 0),
            reporte.total_puntos,
        ]
        totales = [a + b for a, b in zip(totales, fila)]
        ws.append([
            timezone.localtime(reporte.fecha_hora).replace(tzinfo=None),
            reporte.nombre_trabajador,
            reporte.obra.nombre,
            reporte.subetapa.nombre,
            *fila,
            reporte.get_estado_display(),
            (reporte.revisado_por.get_full_name() or reporte.revisado_por.username) if reporte.revisado_por else '',
            reporte.comentario,
            reporte.comentario_revision,
        ])
        ws.cell(row=ws.max_row, column=1).number_format = 'DD-MM-YYYY HH:MM'

    ultima = ws.max_row
    ws.append(['TOTAL', '', '', '', *totales])
    for celda in ws[ws.max_row]:
        celda.font = Font(bold=True)

    anchos = [17, 24, 26, 22, 12, 14, 18, 13, 22, 20, 40, 40]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{ultima}'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    nombre = f'reportes_marvelec_{consultas.hoy_local():%Y%m%d}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{nombre}"'
    wb.save(response)
    return response


# =========================================================
# CENTRO DE NOTIFICACIONES
# =========================================================

def _notif_json(n):
    return {
        'id': n.id,
        'tipo': n.tipo,
        'titulo': n.titulo,
        'cuerpo': n.cuerpo,
        'leida': n.leida,
        'hace': timesince(n.creada).split(',')[0],
        'url_leer': reverse('marvelec_app:notificacion_leer', args=[n.pk]),
    }


@login_required
def notificaciones(request):
    perfil = _get_perfil(request.user)
    pagina = Paginator(Notificacion.objects.filter(usuario=request.user), 30).get_page(request.GET.get('page'))
    plantilla = 'marvelec_app/supervision/notificaciones.html' if perfil.es_supervisor_o_gerencia \
        else 'marvelec_app/trabajador/notificaciones.html'
    return render(request, plantilla, {'pagina': pagina, 'perfil': perfil})


@never_cache
@login_required
def notificaciones_estado(request):
    """JSON para la campanita: cantidad de no leídas y últimas 10."""
    qs = Notificacion.objects.filter(usuario=request.user)
    data = {
        'no_leidas': qs.filter(leida=False).count(),
        'ultimas': [_notif_json(n) for n in qs[:10]],
    }
    perfil = _get_perfil(request.user)
    if perfil.es_supervisor_o_gerencia:
        data['pendientes_revision'] = consultas.reportes_visibles(perfil).filter(
            estado=ReporteAvance.ESTADO_PENDIENTE
        ).exclude(trabajador=request.user).count()
    return JsonResponse(data)


@require_POST
@login_required
def notificacion_leer(request, pk):
    n = get_object_or_404(Notificacion, pk=pk, usuario=request.user)
    if not n.leida:
        n.leida = True
        n.save(update_fields=['leida'])
    destino = n.url or reverse('marvelec_app:notificaciones')
    if _es_fetch(request):
        return JsonResponse({'ok': True, 'url': destino})
    return redirect(destino)


@require_POST
@login_required
def notificaciones_leer_todas(request):
    Notificacion.objects.filter(usuario=request.user, leida=False).update(leida=True)
    if _es_fetch(request):
        return JsonResponse({'ok': True})
    return redirect('marvelec_app:notificaciones')


# =========================================================
# TAREAS PROGRAMADAS (disparadas por GitHub Actions)
# =========================================================

@csrf_exempt  # protegido por el token X-Cron-Token
@require_POST
def tarea_resumen_diario(request):
    token = request.headers.get('X-Cron-Token', '')
    if not settings.CRON_TOKEN or not hmac.compare_digest(token, settings.CRON_TOKEN):
        return JsonResponse({'ok': False, 'mensaje': 'No autorizado'}, status=403)
    from .resumen import enviar_resumen_diario
    resultado = enviar_resumen_diario(consultas.hoy_local())
    return JsonResponse({'ok': True, **resultado})


# =========================================================
# PWA: Service Worker y manifest servidos desde la raíz
# =========================================================

def service_worker(request):
    contenido = render_to_string('marvelec_app/pwa/sw.js', {'version': SW_VERSION})
    response = HttpResponse(contenido, content_type='application/javascript; charset=utf-8')
    response['Cache-Control'] = 'no-cache'
    response['Service-Worker-Allowed'] = '/'
    return response


def manifest(request):
    contenido = render_to_string('marvelec_app/pwa/manifest.webmanifest')
    response = HttpResponse(contenido, content_type='application/manifest+json; charset=utf-8')
    response['Cache-Control'] = 'public, max-age=3600'
    return response
