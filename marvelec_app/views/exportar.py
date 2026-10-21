# marvelec_app/views/exportar.py
"""Exportaciones a Excel para supervisión: reportes, estado de puntos y asistencia mensual."""
from collections import defaultdict
from datetime import date

from django.http import HttpResponse
from django.utils import timezone

from .. import consultas, etapas
from ..models import Asistencia, AvancePunto, Incidencia, Punto
from .comunes import revisor_requerido


def _libro(titulo, encabezados, anchos):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = titulo
    hoja.append(encabezados)
    relleno = PatternFill('solid', fgColor='3A288F')
    for celda in hoja[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = relleno
        celda.alignment = Alignment(vertical='center', wrap_text=True)
    hoja.row_dimensions[1].height = 30
    hoja.freeze_panes = 'A2'
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    return libro, hoja


def _respuesta(libro, nombre):
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{nombre}"'
    libro.save(response)
    return response


def _local(fecha_hora):
    return timezone.localtime(fecha_hora).replace(tzinfo=None) if fecha_hora else None


@revisor_requerido
def exportar_reportes_excel(request):
    """Una fila por punto y etapa registrada, con los filtros de la pantalla."""
    filtros = consultas.leer_filtros(request.GET)
    reportes = consultas.aplicar_filtros(consultas.reportes_visibles(request.perfil), filtros)
    avances = consultas.avances_de(reportes, filtros).select_related(
        'punto__tipo', 'punto__paquete__zona', 'reporte__trabajador', 'reporte__obra', 'reporte__rack',
        'reporte__revisado_por',
    ).order_by('-reporte__fecha_hora', 'reporte_id', 'punto__tipo__sigla', 'punto__numero')

    libro, hoja = _libro('Avance por punto', [
        'Fecha', 'Reporte', 'Trabajador', 'Obra', 'Rack', 'Paquete', 'Zona', 'Punto', 'Etapa', 'Estado del reporte',
        'Revisado por', 'Comentario',
    ], [17, 9, 24, 26, 9, 9, 22, 11, 22, 20, 20, 40])
    for a in avances:
        r = a.reporte
        paquete = a.punto.paquete
        hoja.append([
            _local(r.fecha_hora), r.pk, r.nombre_trabajador, r.obra.nombre, r.rack.nombre,
            paquete.numero if paquete else '', paquete.zona.nombre if paquete and paquete.zona else '',
            a.punto.codigo, a.get_etapa_display(), r.get_estado_display(),
            (r.revisado_por.get_full_name() or r.revisado_por.username) if r.revisado_por else '', r.comentario,
        ])
        hoja.cell(row=hoja.max_row, column=1).number_format = 'DD-MM-YYYY HH:MM'
    hoja.auto_filter.ref = hoja.dimensions
    return _respuesta(libro, f'avance_marvelec_{consultas.hoy_local():%Y%m%d}.xlsx')


@revisor_requerido
def exportar_estado_puntos(request):
    """
    Una fila por punto con la fecha y el trabajador de cada etapa e incidencias abiertas:
    reemplaza la planilla que el supervisor armaba a mano a partir de WhatsApp.
    """
    perfil = request.perfil
    obras = consultas.obras_visibles(perfil)
    obra_id = consultas.parse_int(request.GET.get('obra'))
    if obra_id:
        obras = obras.filter(pk=obra_id)
    puntos = Punto.objects.filter(rack__obra__in=obras).select_related(
        'rack__obra', 'tipo', 'paquete__zona'
    ).order_by('rack__obra__nombre', 'rack__nombre', 'tipo__orden', 'tipo__sigla', 'numero')
    rack_id = consultas.parse_int(request.GET.get('rack'))
    if rack_id:
        puntos = puntos.filter(rack_id=rack_id)

    avances = defaultdict(dict)
    for a in AvancePunto.objects.filter(punto__in=puntos).select_related('reporte__trabajador'):
        avances[a.punto_id][a.etapa] = a
    abiertas = defaultdict(list)
    for inc in Incidencia.objects.filter(punto__in=puntos, estado=Incidencia.ESTADO_ABIERTA).select_related('tipo'):
        abiertas[inc.punto_id].append(inc.tipo.nombre)

    encabezados = ['Obra', 'Rack', 'Paquete', 'Zona', 'Punto', 'Metros']
    for e in etapas.ETAPAS:
        encabezados += [e.nombre, f'{e.corto}: quién']
    encabezados += ['Etapas hechas', 'Incidencia abierta']
    libro, hoja = _libro('Estado de puntos', encabezados,
                         [24, 8, 9, 22, 11, 8] + [12, 18] * len(etapas.ETAPAS) + [10, 28])
    for p in puntos:
        fila = [p.rack.obra.nombre, p.rack.nombre, p.paquete.numero if p.paquete else '',
                p.paquete.zona.nombre if p.paquete and p.paquete.zona else '', p.codigo, p.metros_efectivos]
        for e in etapas.ETAPAS:
            a = avances[p.pk].get(e.clave)
            fila += [timezone.localtime(a.reporte.fecha_hora).date() if a else None,
                     a.reporte.nombre_trabajador if a else '']
        fila += [f'{len(avances[p.pk])}/{len(etapas.ETAPAS)}', ', '.join(abiertas[p.pk])]
        hoja.append(fila)
        for i in range(len(etapas.ETAPAS)):
            hoja.cell(row=hoja.max_row, column=7 + 2 * i).number_format = 'DD-MM-YYYY'
    hoja.auto_filter.ref = hoja.dimensions
    return _respuesta(libro, f'estado_puntos_marvelec_{consultas.hoy_local():%Y%m%d}.xlsx')


@revisor_requerido
def exportar_asistencia(request):
    """Asistencia del mes: días confirmados por trabajador (base para remuneraciones)."""
    hoy = consultas.hoy_local()
    try:
        anio, mes = (int(x) for x in request.GET.get('mes', '').split('-'))
        date(anio, mes, 1)
    except ValueError:
        anio, mes = hoy.year, hoy.month
    obra_id = consultas.obra_fija(request.perfil) or consultas.parse_int(request.GET.get('obra'))
    dias, filas = consultas.asistencia_mensual(request.perfil, anio, mes, obra_id)

    simbolos = {Asistencia.ESTADO_CONFIRMADA: '✓', Asistencia.ESTADO_PENDIENTE: '?', Asistencia.ESTADO_RECHAZADA: '✗'}
    libro, hoja = _libro(
        f'Asistencia {mes:02d}-{anio}',
        ['Trabajador', 'Obra'] + [d.day for d in dias] + ['Confirmados', 'Pendientes', 'Rechazados', 'Sin marca (L-V)'],
        [26, 22] + [4] * len(dias) + [12, 11, 11, 14],
    )
    for f in filas:
        u = f['perfil'].usuario
        sin_marca = sum(1 for d, c in zip(dias, f['celdas']) if c is None and d.weekday() < 5 and d <= hoy)
        hoja.append(
            [u.get_full_name() or u.username, str(f['perfil'].obra_asignada or '')]
            + [simbolos.get(c.estado, '') if c else '' for c in f['celdas']]
            + [f['confirmados'], f['pendientes'], f['rechazados'], sin_marca]
        )
    hoja.append([])
    hoja.append(['✓ confirmada · ? pendiente de confirmar · ✗ rechazada · vacío: no marcó'])
    return _respuesta(libro, f'asistencia_marvelec_{anio}{mes:02d}.xlsx')
